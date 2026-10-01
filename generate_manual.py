from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Title
title = doc.add_heading('Fuel Expense Autopilot', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_paragraph('Official User & Administrator Manual')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].font.size = Pt(14)

doc.add_page_break()

# Table of Contents (Simulated)
doc.add_heading('Table of Contents', level=1)
doc.add_paragraph('1. System Overview')
doc.add_paragraph('2. Field Agent Guide (Mobile App)')
doc.add_paragraph('3. Manager & Approver Guide')
doc.add_paragraph('4. Master Admin Guide (Managing Director)')
doc.add_paragraph('5. Accounts & Finance Guide')

doc.add_page_break()

# 1. System Overview
doc.add_heading('1. System Overview', level=1)
doc.add_paragraph(
    "Fuel Expense Autopilot is a modern, GPS-secured cloud platform designed to track, manage, "
    "and approve employee travel expenses in real-time. It completely replaces legacy Google Forms "
    "and spreadsheets by introducing native GPS watermarking, un-falsifiable camera captures, "
    "and intelligent role-based dashboards."
)

doc.add_heading('System Access URLs', level=2)
doc.add_paragraph('Web Portal (For Managers, Admins, Accounts): https://fuel-expense-autopilot.vercel.app', style='List Bullet')
doc.add_paragraph('Mobile App (For Field Agents): Download the Android APK from the provided secure link.', style='List Bullet')

# 2. Field Agent Guide
doc.add_heading('2. Field Agent Guide (Mobile App)', level=1)
doc.add_paragraph("This section is for employees logging their daily travel and expenses on the road.")

doc.add_heading('Logging In', level=2)
p = doc.add_paragraph('1. Open the Fuel Autopilot app on your Android device.\n')
p.add_run('2. Enter your assigned Email Address and Password.\n')
p.add_run('3. If you forget your password, tap "Forgot Password?" to receive an email with a secure temporary reset code.')

doc.add_heading('Starting Your Day (Check-In)', level=2)
p = doc.add_paragraph('1. On the home dashboard, tap the blue ')
p.add_run('START TRIP').bold = True
p.add_run(' button.\n2. Enter your starting vehicle Odometer reading.\n3. Tap the Camera box. The app will automatically capture your GPS location and time, and securely stamp it onto your dashboard photo.\n4. Click ')
p.add_run('Confirm & Submit').bold = True
p.add_run('. Your trip is now active!')

doc.add_heading('Adding an Expense (Fuel, Tolls, Servicing)', level=2)
p = doc.add_paragraph('1. During your trip, tap the green ')
p.add_run('Add Bill').bold = True
p.add_run(' button.\n2. Select the Expense Type (Fuel, Servicing, Spare Parts, Toll/Parking, Other).\n3. Enter the total Amount in Rupees.\n4. Take a live photo of the receipt. The system will auto-stamp your GPS coordinates to prove you were at the fuel station/shop.\n5. Tap Submit. This instantly sends the bill to your Manager.')

doc.add_heading('Ending Your Day (Check-Out)', level=2)
p = doc.add_paragraph('1. At the end of your shift, tap the red ')
p.add_run('END TRIP').bold = True
p.add_run(' button.\n2. Enter your final Odometer reading.\n3. Type out the exact Route and Locations you visited today (e.g., GS Road, Panbazar).\n4. Take a final photo of your dashboard to verify the Odometer reading.')

# 3. Manager & Approver Guide
doc.add_heading('3. Manager & Team Leader Guide', level=1)
doc.add_paragraph("Managers use the Web Portal on their computer to monitor teams and approve claims.")

doc.add_heading('Pending Approvals', level=2)
p = doc.add_paragraph('1. Click on ')
p.add_run('Pending Approvals').bold = True
p.add_run(' in the left sidebar.\n2. You will see a list of all expenses submitted by your agents in real-time.\n3. Review the variance (e.g., if claimed distance heavily mismatches OSRM calculated distance, it will highlight in red).\n4. Click the checkmark to Approve, the X to Reject, or the Flag icon to investigate further.')

# 4. Master Admin Guide
doc.add_heading('4. Master Admin Guide (Managing Director)', level=1)
doc.add_paragraph("The Master Admin (admin@dhanpurna.net) has exclusive access to system configuration and employee management.")

doc.add_heading('User Management', level=2)
p = doc.add_paragraph('Only the Master Admin can access the ')
p.add_run('User Management').bold = True
p.add_run(' tab. Here you can:\n')
doc.add_paragraph('Add New Users: Create accounts for new hires, assign them an employee code, set their initial password, and select their Role (Field Agent, Manager, Accounts, etc.)', style='List Bullet')
doc.add_paragraph('Edit Users: Click the Pencil icon next to an employee to change their department, update their email, or reset their password.', style='List Bullet')
doc.add_paragraph('Delete/Deactivate Users: Click the Trash Can icon to permanently remove an employee from the system.', style='List Bullet')

# 5. Accounts Guide
doc.add_heading('5. Accounts & Finance Guide', level=1)
doc.add_paragraph("Finance users have access to fully approved claims ready for payout.")
doc.add_paragraph('1. Click Reconciliation Queue to view claims approved by managers.', style='List Bullet')
doc.add_paragraph('2. Export data into CSV/Excel format for Tally or direct bank payouts using the Export Center.', style='List Bullet')

doc.save(r"C:\Users\MIS\Desktop\Fuel_Autopilot_User_Manual.docx")
